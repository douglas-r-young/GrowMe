"""export_pptx writes a valid pptx with one slide per Slide and notes attached."""
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE

from growme.decks.pptx import export_pptx
from growme.schemas import SessionDeck, Slide


def _slide(layout, blocks=None, **kw):
    blocks = blocks or {}
    title = blocks.get("title") or blocks.get("behavior_name") or "Untitled"
    return Slide(
        title=str(title),
        body_md=str(blocks.get("subtitle") or blocks.get("prompt") or ""),
        layout=layout,
        blocks=blocks,
        speaker_notes="Talk track for this slide. " * 8,
        **kw,
    )


def _deck() -> SessionDeck:
    return SessionDeck(
        title="Photon DB · PIC Mastery",
        behavior_ids=["a", "b", "c"],
        slides=[
            _slide("cover", {"eyebrow": "GrowMe", "title": "Welcome", "subtitle": "lo"}),
            _slide("poll_qr", {"title": "Pre QR", "body": "scan", "caption": "cap"},
                   kind="poll_qr",
                   qr_url="http://localhost:8501/?assessment=u&kind=pre",
                   qr_caption="cap"),
            _slide("section_divider", {"number": "01", "behavior_name": "Quantify pain", "promise": "p"}),
            _slide("teach", {"eyebrow": "SPIN", "title": "Quantify pain",
                             "bullets": ["a", "b", "c"], "citation": "src"}),
            _slide("example", {"title": "Before/After",
                               "before_body": "before", "after_body": "after"}),
            _slide("activity", {"eyebrow": "Try it · 5 min", "title": "Practice",
                                "prompt": "do this", "sub_prompts": ["x", "y"],
                                "timer_hint": "5 min"}),
            _slide("poll_qr", {"title": "Post QR", "body": "scan", "caption": "cap"},
                   kind="poll_qr",
                   qr_url="http://localhost:8501/?assessment=u&kind=post",
                   qr_caption="cap"),
            _slide("close", {"title": "Thanks",
                             "commitment_recap": "you committed",
                             "next_step": "nudge next week"},
                   kind="close"),
        ],
        facilitator_guide_md="# g",
        pre_qr_url="http://localhost:8501/?assessment=u&kind=pre",
        post_qr_url="http://localhost:8501/?assessment=u&kind=post",
    )


def test_export_pptx_writes_valid_pptx_with_correct_slide_count(tmp_path: Path):
    out = export_pptx(_deck(), tmp_path / "deck.pptx")
    assert out.exists()
    assert out.read_bytes()[:2] == b"PK"

    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        slide_xmls = [n for n in names if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
        notes_xmls = [n for n in names if n.startswith("ppt/notesSlides/notesSlide") and n.endswith(".xml")]
    assert len(slide_xmls) == 8
    # Every slide should have an attached notes slide (we set speaker_notes on each).
    assert len(notes_xmls) == 8


def _shape_with_text(slide, needle: str):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        if needle in sh.text_frame.text:
            return sh
    return None


def test_section_divider_number_sits_above_title(tmp_path: Path):
    """The chapter numeral must end before the title begins — no vertical overlap."""
    out = export_pptx(_deck(), tmp_path / "deck.pptx")
    prs = Presentation(str(out))
    # slide index 2 in our fixture is the section_divider.
    sd = prs.slides[2]

    num = _shape_with_text(sd, "01")
    title = _shape_with_text(sd, "Quantify pain")
    assert num is not None and title is not None
    # ghost number's bottom edge must be at or above the title's top edge
    assert num.top + num.height <= title.top, (
        f"section divider numeral overlaps title: "
        f"num bottom={num.top + num.height}, title top={title.top}"
    )


def test_example_does_not_render_pull_quote_on_slide(tmp_path: Path):
    """pull_quote is routed to speaker notes by the builder; the renderer must
    never put it on the slide even if the field is present in blocks."""
    deck = SessionDeck(
        title="Test",
        behavior_ids=["a", "b", "c"],
        slides=[
            _slide("example", {
                "title": "Before/After",
                "before_body": "before",
                "after_body": "after",
                "pull_quote": "This is a unique pull-quote sentinel string.",
            }),
        ],
        facilitator_guide_md="# g",
        pre_qr_url="http://x", post_qr_url="http://x",
    )
    out = export_pptx(deck, tmp_path / "deck.pptx")
    prs = Presentation(str(out))
    for sh in prs.slides[0].shapes:
        if sh.has_text_frame:
            assert "sentinel" not in sh.text_frame.text, (
                "pull_quote leaked onto the example slide"
            )


def test_close_body_textboxes_use_text_to_fit_shape(tmp_path: Path):
    """The recap and next-step text frames must shrink-to-fit so a worst-case
    LLM response can't overflow into the next visual zone."""
    out = export_pptx(_deck(), tmp_path / "deck.pptx")
    prs = Presentation(str(out))
    close = prs.slides[7]  # last slide in fixture

    matches = [
        sh for sh in close.shapes
        if sh.has_text_frame and (
            "you committed" in sh.text_frame.text
            or "nudge next week" in sh.text_frame.text
        )
    ]
    assert len(matches) == 2, f"expected commitment + next_step textboxes, got {len(matches)}"
    for sh in matches:
        assert sh.text_frame.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
